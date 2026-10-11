"""Exercise real Qt interactions, including the broken search binding."""
import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

pytest.importorskip("PySide6")
from PySide6.QtCore import QPoint, QPointF, QSettings, Qt, QTimer
from PySide6.QtGui import QFontDatabase, QImage, QPalette, QWheelEvent
from PySide6.QtPdf import QPdfDocument, QPdfDocumentRenderOptions
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFileDialog, QInputDialog, QMessageBox, QTabBar

from ape_pdf_studio.ui.main_window import MainWindow
from ape_pdf_studio.ui.theme import DEFAULT_THEME, THEMES, apply_theme, theme_colors, theme_color
from ape_pdf_studio.ui.settings_dialog import SettingsDialog
from ape_pdf_studio.ui.exports import DocxExportJob, export_page_image
from ape_pdf_studio.core.editor import PDFOperationError
from ape_pdf_studio.ui.properties_dialog import PropertiesDialog
from ape_pdf_studio.ui.content_dialogs import NoteDialog, TextEditDialog


@pytest.fixture(scope="session")
def app(tmp_path_factory):
    application = QApplication.instance() or QApplication([])
    QSettings.setDefaultFormat(QSettings.Format.IniFormat)
    QSettings.setPath(QSettings.Format.IniFormat, QSettings.Scope.UserScope, str(tmp_path_factory.mktemp("settings")))
    # The Windows offscreen plugin does not enumerate system fonts.
    font = Path("C:/Windows/Fonts/segoeui.ttf")
    if font.exists():
        QFontDatabase.addApplicationFont(str(font))
    apply_theme(application)
    return application


@pytest.fixture
def text_pdf(tmp_path):
    writer = PdfWriter()
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    font_ref = writer._add_object(font)
    for number in range(3):
        page = writer.add_blank_page(width=400, height=600)
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 16 Tf 48 500 Td (Alpha beta PDF page {number + 1}) Tj 0 -32 Td (Second line for selection) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_outline_item("First page", 0)
    path = tmp_path / "text.pdf"
    writer.write(path)
    return path


@pytest.fixture
def window(app, text_pdf, monkeypatch):
    errors = []
    monkeypatch.setattr(QMessageBox, "critical", lambda *args: errors.append(args[2]))
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: errors.append(args[2]))
    widget = MainWindow()
    widget.show()
    widget.open_pdf(text_pdf)
    QTest.qWait(120)
    assert widget.current is not None, errors
    yield widget
    for index in reversed(range(widget.tabs.count())):
        tab = widget.tabs.widget(index)
        tab.close_document()
        widget.tabs.removeTab(index)
        tab.deleteLater()
    widget.close()
    QTest.qWait(20)
    assert not errors


def wait_for(predicate, timeout=2500):
    for _ in range(timeout // 25):
        if predicate():
            return
        QTest.qWait(25)
    assert predicate()


def test_search_typing_results_navigation_and_clear(window):
    tab = window.current
    window.focus_search()
    QTest.mouseClick(window.search, Qt.MouseButton.LeftButton)
    QTest.keyClicks(window.search, "PDF")
    wait_for(lambda: tab.search_model.count() == 3)
    assert tab._selected_search == 0
    assert tab.viewer.currentSearchResultIndex() == 0
    assert "1 di 3" in window.search_status.text()
    window._search_step(1)
    assert tab.page == 1
    assert window.page_spin.value() == 2
    window._search_step(-1)
    window._search_step(-1)
    assert tab.page == 2
    window.search.clear()
    wait_for(lambda: tab.search_model.count() == 0)
    assert tab.viewer.currentSearchResultIndex() == -1
    assert not window.a_find_next.isEnabled()


def test_drag_copy_context_menu_markup_and_history(window, app, text_pdf):
    tab = window.current
    view = tab.viewer
    bounds = tab.pdf.getAllText(0).bounds()[0].boundingRect()
    transform = view.page_transform(0)
    start = transform.map(QPointF(bounds.left() + 1, bounds.center().y())).toPoint()
    end = transform.map(QPointF(bounds.right() - 1, bounds.center().y())).toPoint()
    QTest.mousePress(view.viewport(), Qt.MouseButton.LeftButton, pos=start)
    QTest.mouseMove(view.viewport(), end, 20)
    QTest.mouseRelease(view.viewport(), Qt.MouseButton.LeftButton, pos=end)
    assert view.has_selection
    assert "Alpha beta PDF" in view.selection.text()
    QTest.keyClick(view, Qt.Key.Key_C, Qt.KeyboardModifier.ControlModifier)
    assert app.clipboard().text() == view.selection.text()
    menu = view.make_context_menu()
    assert menu.actions()[0].isEnabled()
    assert menu.actions()[2].isEnabled()
    menu.actions()[2].trigger()
    wait_for(lambda: tab.editor.dirty and tab.count == 3)
    reader = PdfReader(tab.editor.path)
    assert reader.pages[0]["/Annots"][0].get_object()["/Subtype"] == "/Highlight"
    reader.close()
    assert not PdfReader(text_pdf).pages[0].get("/Annots")
    window.undo()
    assert not tab.editor.dirty
    window.redo()
    assert tab.editor.dirty
    # Branching while old render jobs exist must not hit Windows file locks.
    window.undo()
    window._edit("rotate", 0, 90)
    assert not tab.editor.can_redo
    menu.deleteLater()


def test_select_all_copy_and_underline_save(window, app, tmp_path, monkeypatch):
    tab = window.current
    QTest.keyClick(tab.viewer, Qt.Key.Key_A, Qt.KeyboardModifier.ControlModifier)
    assert tab.viewer.has_selection
    window.a_copy.trigger()
    assert "Second line" in app.clipboard().text()
    window.a_underline.trigger()
    destination = tmp_path / "annotated"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(destination), ""))
    window.save()
    saved = destination.with_suffix(".pdf")
    assert saved.exists()
    assert window.toast.isVisible() and window.toast.title.text() == "PDF salvato"
    assert window.toast.filename.text() == saved.name
    assert not tab.editor.dirty
    reader = PdfReader(saved)
    assert reader.pages[0]["/Annots"][0].get_object()["/Subtype"] == "/Underline"
    reader.close()
    window._edit("rotate", 1, 90)
    window.save()
    assert not tab.editor.dirty
    assert PdfReader(saved).pages[1].rotation == 90
    assert window.toast.isVisible() and window.toast.filename.text() == saved.name


def test_page_controls_and_zoom_do_not_disappear(window):
    tab = window.current
    assert not window.a_prev.isEnabled()
    assert not window.a_up.isEnabled()
    window.page_spin.setValue(3)
    assert tab.page == 2
    assert not window.a_next.isEnabled()
    assert not window.a_down.isEnabled()
    tab.zoom_in()
    assert tab.viewer.zoomFactor() > 0
    tab.set_zoom(150)
    assert tab.viewer.zoomFactor() == 1.5
    tab.jump_to_page(0)
    assert tab.viewer.zoomFactor() == 1.5
    tab.fit_width()
    previous = tab.effective_zoom()
    tab.zoom_out()
    assert tab.viewer.zoomFactor() == pytest.approx(previous / 1.2)


def test_double_click_selects_word_and_reverse_drag(window):
    tab = window.current
    view = tab.viewer
    first_word = tab.pdf.getSelectionAtIndex(0, 0, 5).boundingRectangle()
    point = view.page_transform(0).map(first_word.center()).toPoint()
    QTest.mouseDClick(view.viewport(), Qt.MouseButton.LeftButton, pos=point)
    assert view.selection.text() == "Alpha"
    bounds = tab.pdf.getAllText(0).bounds()[0].boundingRect()
    transform = view.page_transform(0)
    start = transform.map(QPointF(bounds.right() - 1, bounds.center().y())).toPoint()
    end = transform.map(QPointF(bounds.left() + 1, bounds.center().y())).toPoint()
    QTest.mousePress(view.viewport(), Qt.MouseButton.LeftButton, pos=start)
    QTest.mouseMove(view.viewport(), end, 20)
    QTest.mouseRelease(view.viewport(), Qt.MouseButton.LeftButton, pos=end)
    assert "Alpha beta PDF" in view.selection.text()


def test_search_no_match_and_empty_page_selection(window, tmp_path, app):
    window.search.setText("does-not-exist")
    QTest.qWait(300)
    assert "Nessun risultato" in window.search_status.text()
    blank = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=400, height=600)
    writer.write(blank)
    window.open_pdf(blank)
    tab = window.current
    tab.viewer.select_page_text()
    assert not tab.viewer.has_selection
    assert not window.a_copy.isEnabled()
    assert "OCR" in window.statusBar().currentMessage()


def test_save_over_original_retains_live_view_and_history(window, text_pdf, monkeypatch):
    original = text_pdf.read_bytes()
    window._edit("rotate", 0, 90)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(text_pdf), ""))
    monkeypatch.setattr(window, "_confirm_overwrite", lambda _: True)
    window.save_as()
    assert not window.current.editor.dirty
    assert PdfReader(text_pdf).pages[0].rotation == 90
    window.undo()
    assert window.current.editor.path.read_bytes() == original
    assert window.current.count == 3


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_markup_rendered_on_rotated_cropped_page(window, tmp_path, rotation):
    tab = window.current
    window._edit("rotate", 0, rotation or 180)
    if rotation == 0:
        window.undo()
    # Add nonzero crop offsets to validate visual-to-PDF coordinates.
    reader = PdfReader(tab.editor.path)
    writer = PdfWriter()
    writer.append(reader)
    writer.pages[0].cropbox.lower_left = (20, 30)
    writer.pages[0].cropbox.upper_right = (380, 550)
    cropped = tmp_path / "cropped.pdf"
    writer.write(cropped)
    reader.close()
    window.open_pdf(cropped)
    tab = window.current
    view = tab.viewer
    bounds = tab.pdf.getAllText(0).bounds()[0].boundingRect()
    size = tab.pdf.pagePointSize(0).toSize()
    options = QPdfDocumentRenderOptions()
    options.setRenderFlags(QPdfDocumentRenderOptions.RenderFlag.Annotations)
    before = tab.pdf.render(0, size, options)
    view.selection = tab.pdf.getAllText(0)
    view.selection_page = 0
    window.markup("highlight")
    after = tab.pdf.render(0, size, options)
    changed = 0
    rect = bounds.toAlignedRect().intersected(after.rect())
    for y in range(rect.top(), rect.bottom() + 1):
        for x in range(rect.left(), rect.right() + 1):
            if before.pixel(x, y) != after.pixel(x, y):
                changed += 1
    assert changed > 30


def test_closing_dirty_tab_can_save_and_cancel(window, tmp_path, monkeypatch):
    window._edit("rotate", 0, 90)
    monkeypatch.setattr(QMessageBox, "question", lambda *args: QMessageBox.StandardButton.Save)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: ("", ""))
    close = window.tabs.tabBar().tabButton(0, QTabBar.ButtonPosition.RightSide)
    QTest.mouseClick(close, Qt.MouseButton.LeftButton)
    assert window.tabs.count() == 1
    saved = tmp_path / "close-saved.pdf"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(saved), ""))
    QTest.mouseClick(close, Qt.MouseButton.LeftButton)
    assert window.tabs.count() == 0
    assert saved.exists()


def test_tab_close_after_reordering_keeps_other_document_active(window, text_pdf, tmp_path, monkeypatch):
    first = window.current
    window._edit("rotate", 0, 90)
    bar = window.tabs.tabBar()
    close_first = bar.tabButton(0, QTabBar.ButtonPosition.RightSide)
    other = tmp_path / "A long document filename with an identifiable ending 2026.pdf"
    other.write_bytes(text_pdf.read_bytes())
    window.open_pdf(other)
    second = window.current
    bar.moveTab(0, 1)
    assert window.tabs.widget(1) is first and window.tabs.widget(0) is second
    assert window.tabs.tabText(1).endswith(" *")
    assert window.tabs.tabToolTip(0) == str(other.resolve())
    monkeypatch.setattr(QMessageBox, "question", lambda *args: QMessageBox.StandardButton.Cancel)
    QTest.mouseClick(close_first, Qt.MouseButton.LeftButton)
    assert window.tabs.count() == 2 and window.current is second
    monkeypatch.setattr(QMessageBox, "question", lambda *args: QMessageBox.StandardButton.Discard)
    QTest.mouseClick(close_first, Qt.MouseButton.LeftButton)
    assert window.tabs.count() == 1 and window.current is second
    assert second.pdf.status() == QPdfDocument.Status.Ready


def test_search_is_per_tab_when_switching_before_debounce(window, text_pdf, tmp_path):
    first = window.current
    window.search.setText("Alpha")
    other = tmp_path / "other.pdf"
    other.write_bytes(text_pdf.read_bytes())
    window.open_pdf(other)
    assert window.search.text() == ""
    window.search.setText("Second")
    wait_for(lambda: window.current.search_model.count() == 3)
    window.tabs.setCurrentIndex(0)
    assert window.search.text() == "Alpha"
    wait_for(lambda: first.search_model.count() == 3)
    assert first.search_model.searchString() == "Alpha"


def test_recent_documents_and_settings_are_persisted(window, text_pdf):
    window.settings.sync()
    assert window.settings.format() == QSettings.Format.IniFormat
    restored = MainWindow()
    assert str(text_pdf.resolve()) in restored.recent_files
    restored.close()


@pytest.mark.parametrize("modifier", [Qt.KeyboardModifier.ControlModifier, Qt.KeyboardModifier.MetaModifier])
def test_wheel_zoom_at_pointer_on_scrolled_page(window, modifier):
    view = window.current.viewer
    # Start with a page wider than the viewport: centering a smaller page
    # necessarily clamps horizontal panning at its scroll limits.
    window.current.set_zoom(250)
    window.current.jump_to_page(1)
    point = QPointF(130, 100)
    position = view.page_transform(1).map(point)
    event = QWheelEvent(position, QPointF(view.viewport().mapToGlobal(position.toPoint())),
                        QPoint(), QPoint(0, 120), Qt.MouseButton.NoButton, modifier,
                        Qt.ScrollPhase.NoScrollPhase, False)
    QApplication.sendEvent(view.viewport(), event)
    assert view.zoomFactor() == pytest.approx(3.0)
    after = view.page_transform(1).map(point)
    assert abs(after.x() - position.x()) <= 1
    assert abs(after.y() - position.y()) <= 1
    assert window.zoom_combo.currentText() == "300%"


def test_normal_wheel_scroll_and_pixel_zoom_limits(window):
    view = window.current.viewer
    window.current.set_zoom(200)
    position = QPointF(view.viewport().rect().center())
    def wheel(angle, pixels=0, modifier=Qt.KeyboardModifier.NoModifier):
        event = QWheelEvent(position, position, QPoint(0, pixels), QPoint(0, angle),
                            Qt.MouseButton.NoButton, modifier, Qt.ScrollPhase.NoScrollPhase, False)
        QApplication.sendEvent(view.viewport(), event)
    wheel(-120)
    assert view.zoomFactor() == 2
    assert view.verticalScrollBar().value() > 0
    wheel(0, 15, Qt.KeyboardModifier.ControlModifier)
    assert view.zoomFactor() > 2
    window.current.set_zoom(500)
    wheel(120, modifier=Qt.KeyboardModifier.MetaModifier)
    assert view.zoomFactor() == 5
    window.current.set_zoom(25)
    wheel(-120, modifier=Qt.KeyboardModifier.ControlModifier)
    assert view.zoomFactor() == 0.25


@pytest.mark.parametrize("vector", [True, False])
def test_middle_button_pan_moves_both_axes_and_preserves_selection(window, vector):
    view = window.current.viewer
    view.set_vector_enabled(vector)
    window.current.set_zoom(500)
    view.select_page_text(0)
    selection = view.selection.text()
    horizontal, vertical = view.horizontalScrollBar(), view.verticalScrollBar()
    horizontal.setValue(horizontal.maximum() // 2)
    vertical.setValue(vertical.maximum() // 3)
    before = horizontal.value(), vertical.value()
    original_cursor = view.viewport().cursor().shape()
    start = QPoint(200, 200)
    end = start + QPoint(60, 40)
    QTest.mousePress(view.viewport(), Qt.MouseButton.MiddleButton, pos=start)
    assert view.viewport().cursor().shape() == Qt.CursorShape.ClosedHandCursor
    QTest.mouseMove(view.viewport(), end, 20)
    assert (horizontal.value(), vertical.value()) == (before[0] - 60, before[1] - 40)
    assert view.selection.text() == selection and not window.current.editor.dirty
    QTest.mouseRelease(view.viewport(), Qt.MouseButton.MiddleButton, pos=end)
    assert view.viewport().cursor().shape() == original_cursor
    after = horizontal.value(), vertical.value()
    QTest.mouseMove(view.viewport(), end + QPoint(20, 20), 20)
    assert (horizontal.value(), vertical.value()) == after
    assert view.zoomFactor() == 5


def test_middle_button_pan_clamps_and_can_reverse_at_page_boundary(window):
    view = window.current.viewer
    window.current.set_zoom(500)
    horizontal, vertical = view.horizontalScrollBar(), view.verticalScrollBar()
    horizontal.setValue(0)
    vertical.setValue(0)
    QTest.mousePress(view.viewport(), Qt.MouseButton.MiddleButton, pos=QPoint(200, 200))
    QTest.mouseMove(view.viewport(), QPoint(260, 240), 20)
    assert (horizontal.value(), vertical.value()) == (0, 0)
    QTest.mouseMove(view.viewport(), QPoint(230, 220), 20)
    assert (horizontal.value(), vertical.value()) == (30, 20)
    QTest.mouseRelease(view.viewport(), Qt.MouseButton.MiddleButton, pos=QPoint(230, 220))


def test_middle_button_pan_stops_on_escape_and_focus_loss(window):
    view = window.current.viewer
    window.current.set_zoom(500)
    view.select_page_text(0)
    selected = view.selection.text()
    QTest.mousePress(view.viewport(), Qt.MouseButton.MiddleButton, pos=QPoint(200, 200))
    QTest.keyClick(view, Qt.Key.Key_Escape)
    assert view._pan_position is None and view.selection.text() == selected
    QTest.mouseRelease(view.viewport(), Qt.MouseButton.MiddleButton, pos=QPoint(200, 200))
    QTest.mousePress(view.viewport(), Qt.MouseButton.MiddleButton, pos=QPoint(200, 200))
    assert view._pan_position is not None
    window.zoom_combo.setFocus()
    assert view._pan_position is None
    assert view.viewport().cursor().shape() != Qt.CursorShape.ClosedHandCursor
    QTest.mouseRelease(view.viewport(), Qt.MouseButton.MiddleButton, pos=QPoint(200, 200))


def test_compact_workspace_and_collapsible_search(window):
    assert not window.search_toolbar.isVisible()
    assert window.navigation_control.isVisible()
    window.resize(900, 620)
    QTest.qWait(60)
    height = window.current.viewer.viewport().height()
    assert height >= window.height() - 160 - window.navigation_toolbar.height()
    window.focus_search()
    assert window.search_toolbar.isVisible()
    window.search.setText("Alpha")
    wait_for(lambda: window.current.search_model.count() == 3)
    QTest.keyClick(window.search, Qt.Key.Key_Escape)
    assert not window.search_toolbar.isVisible()
    assert window.current.sidebar.currentWidget() is window.current.thumbs
    wait_for(lambda: window.current.search_model.count() == 0)
    window.close_tab(0)
    assert not window.navigation_control.isVisible()
    assert not window.annotate_control.isVisible()


def test_thumbnail_scroll_click_and_reload_for_long_pdf(window, tmp_path):
    writer = PdfWriter()
    for number in range(95):
        page = writer.add_blank_page(width=400, height=600)
        stream = DecodedStreamObject()
        stream.set_data(f"{number / 100:.2f} 0.5 0.8 rg 0 0 400 600 re f".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    path = tmp_path / "long.pdf"
    writer.write(path)
    window.open_pdf(path)
    tab = window.current
    wait_for(lambda: not tab.thumbs.item(0).icon().isNull())
    tab.thumbs.verticalScrollBar().setValue(tab.thumbs.verticalScrollBar().maximum())
    wait_for(lambda: not tab.thumbs.item(94).icon().isNull())
    assert min(tab.thumbs.visible_pages()) > 85
    assert tab.page == 0  # Browsing thumbnails alone does not move the document.
    rect = tab.thumbs.visualItemRect(tab.thumbs.item(94))
    QTest.mouseClick(tab.thumbs.viewport(), Qt.MouseButton.LeftButton, pos=rect.center())
    assert tab.page == 94
    assert window.page_spin.value() == 95
    window._edit("rotate", 94, 90)
    wait_for(lambda: not tab.thumbs.item(94).icon().isNull())
    size = tab.thumbs.item(94).icon().availableSizes()[0]
    assert size.width() > size.height()
    window.undo()
    wait_for(lambda: not tab.thumbs.item(94).icon().isNull())
    size = tab.thumbs.item(94).icon().availableSizes()[0]
    assert size.height() > size.width()
    tab.jump_to_page(0)
    wait_for(lambda: not tab.thumbs.item(0).icon().isNull())
    assert tab.thumbs.currentRow() == 0
    assert tab.thumbs.verticalScrollBar().value() < 200
    assert len(tab.thumbs.cache) <= 72


@pytest.mark.parametrize("restricted", [False, True])
def test_document_notices_follow_original_after_edit_save_undo(window, tmp_path, restricted):
    from test_document_info import add_pdfa, add_signature
    assert not window.current.notice.isVisible()
    writer = PdfWriter()
    writer.add_blank_page(400, 600)
    if not restricted:
        add_pdfa(writer)
    add_signature(writer)
    if restricted:
        writer.encrypt("", owner_password="owner-secret")
    source = tmp_path / "declared-and-signed.pdf"
    writer.write(source)
    window.open_pdf(source)
    tab = window.current
    assert tab.notice.isVisible()
    QTest.qWait(25)
    assert tab.notice.height() <= 64
    if not restricted:
        assert "PDF/A-2b dichiarato" in tab.notice.text.text()
    assert "1 firma digitale" in tab.notice.text.text()
    assert "validità non verificata" in tab.notice.text.text()
    if restricted:
        assert "Modifiche disabilitate" in tab.notice.text.text()
        assert "sola lettura" in tab.notice.text.text()
        assert not window.a_rotate.isEnabled()
        return
    assert "Modifiche disabilitate" not in tab.notice.text.text()
    window._edit("rotate", 0, 90)
    assert "PDF modificato" in tab.notice.text.text()
    assert "conformità da verificare" in tab.notice.text.text()
    assert tab.notice.text.palette().color(QPalette.ColorRole.WindowText).name() == "#392807"
    tab.editor.save_as(tmp_path / "modified.pdf")
    tab.notice.refresh()
    assert "PDF modificato" in tab.notice.text.text()
    assert not tab.editor.dirty
    window.undo()
    assert "PDF modificato" not in tab.notice.text.text()
    assert "dichiarato" in tab.notice.text.text()


@pytest.mark.parametrize("extension", [".png", ".jpg"])
def test_export_current_page_image_with_rotation_markup_and_dpi(window, tmp_path, extension):
    tab = window.current
    tab.jump_to_page(1)
    tab.apply_edit("add_text_markup", 1, [(48, 80, 230, 110)], "highlight")
    tab.apply_edit("rotate", 1, 90)
    wait_for(lambda: tab.pdf.status() == QPdfDocument.Status.Ready)
    output = tmp_path / ("page" + extension)
    export_page_image(tab.pdf, 1, output, 150)
    image = QImage(str(output))
    assert not image.isNull()
    assert (image.width(), image.height()) == (1250, 834)
    assert abs(image.dotsPerMeterX() - round(150 / 0.0254)) < 5
    assert image.pixelColor(2, 2).red() > 245  # White paper, including JPEG.
    colors = [image.pixelColor(x, y) for x in range(900, 1110, 7) for y in range(100, 480, 7)]
    assert any(c.red() > 200 and c.green() > 180 and c.blue() < 100 for c in colors)
    assert tab.editor.dirty  # Export does not mark unsaved PDF edits as saved.


def test_export_image_dialog_uses_current_page_and_jpeg_extension(window, tmp_path, monkeypatch):
    window.current.jump_to_page(2)
    monkeypatch.setattr(QInputDialog, "getItem", lambda *_args, **_kwargs: ("300 dpi", True))
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *_args, **_kwargs: (str(tmp_path / "third"), "Immagine JPEG (*.jpg *.jpeg)"))
    window.export_image()
    image = QImage(str(tmp_path / "third.jpg"))
    assert (image.width(), image.height()) == (1667, 2500)
    assert not window.current.editor.dirty
    assert window.toast.isVisible() and window.toast.title.text() == "Immagine esportata"
    assert window.toast.filename.text() == "third.jpg"


def test_image_export_rejects_pdf_extension_and_excessive_resolution(window, tmp_path):
    tab = window.current
    source_bytes = tab.editor.source.read_bytes()
    with pytest.raises(PDFOperationError, match="PNG"):
        export_page_image(tab.pdf, 0, tab.editor.source)
    assert tab.editor.source.read_bytes() == source_bytes
    with pytest.raises(PDFOperationError, match="pronta"):
        export_page_image(tab.pdf, 20, tmp_path / "bad.png")
    writer = PdfWriter()
    writer.add_blank_page(width=3000, height=3000)
    huge = tmp_path / "huge.pdf"
    writer.write(huge)
    window.open_pdf(huge)
    wait_for(lambda: window.current.pdf.status() == QPdfDocument.Status.Ready)
    with pytest.raises(PDFOperationError, match="40 megapixel"):
        export_page_image(window.current.pdf, 0, tmp_path / "huge.png", 600)
    assert not (tmp_path / "huge.png").exists()


def test_word_export_gui_converts_all_pages_and_preserves_dirty_state(window, tmp_path, monkeypatch):
    from docx import Document

    window.current.apply_edit("add_note", 0, "Pending edit")
    original_revision = window.current.editor.path
    original = window.current.editor.source.read_bytes()
    messages = []
    monkeypatch.setattr(QMessageBox, "information", lambda *args: messages.append(args[2]))
    output = tmp_path / "whole_document"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *_args, **_kwargs: (str(output), ""))
    window.export_docx()
    assert window._export_job is not None
    assert not window.toast.isVisible()
    assert not window.a_export_docx.isEnabled()
    wait_for(lambda: window._export_job is None, timeout=15000)
    exported = Document(output.with_suffix(".docx"))
    text = "\n".join(p.text for p in exported.paragraphs)
    for number in range(1, 4):
        assert f"PDF page {number}" in text
    assert len(exported.sections) == 3
    assert not messages
    assert window.toast.isVisible() and window.toast.title.text() == "Documento Word esportato"
    assert "3 pagine" in window.toast.detail.text() and "impaginazione" in window.toast.detail.text()
    assert window.toast.filename.text() == "whole_document.docx"
    assert window.current.editor.path == original_revision and window.current.editor.dirty
    assert window.current.editor.source.read_bytes() == original
    assert window.a_export_docx.isEnabled()


def test_word_job_snapshot_survives_source_closure(window, tmp_path):
    from docx import Document

    tab = window.current
    output = tmp_path / "snapshot.docx"
    job = DocxExportJob(tab.editor.path, output, window)
    work = job.snapshot.parent
    completed, errors = [], []
    job.finished.connect(lambda: completed.append(True))
    job.failed.connect(errors.append)
    # A separate export snapshot has no dependency on the tab's temp directory.
    index = window.tabs.currentIndex()
    tab.close_document()
    window.tabs.removeTab(index)
    tab.deleteLater()
    job.start()
    wait_for(lambda: completed, timeout=15000)
    assert not errors
    assert len(Document(output).sections) == 3
    assert not work.exists()
    job.deleteLater()


def test_cancelled_word_job_keeps_destination_and_cleans_snapshot(window, tmp_path):
    output = tmp_path / "existing.docx"
    output.write_bytes(b"existing file")
    job = DocxExportJob(window.current.editor.path, output, window)
    work = job.snapshot.parent
    cancelled, completed, errors = [], [], []
    job.cancelled.connect(lambda: cancelled.append(True))
    job.finished.connect(lambda: completed.append(True))
    job.failed.connect(errors.append)
    job.start()
    job.cancel()
    wait_for(lambda: completed, timeout=15000)
    assert cancelled and not errors
    assert output.read_bytes() == b"existing file"
    assert not work.exists()
    job.deleteLater()


def test_metadata_dialog_edit_apply_undo_and_save(window, tmp_path):
    tab = window.current
    tab.jump_to_page(1)
    dialog = PropertiesDialog(tab.editor, lambda values: tab.apply_edit("update_metadata", values), window)
    dialog.show()
    assert not dialog.apply_button.isEnabled()
    title = dialog.fields["/Title"]
    QTest.mouseClick(title, Qt.MouseButton.LeftButton)
    QTest.keyClicks(title, "Edited document title")
    assert dialog.apply_button.isEnabled()
    QTest.mouseClick(dialog.apply_button, Qt.MouseButton.LeftButton)
    assert not dialog.apply_button.isEnabled()
    assert "Metadati aggiornati" in dialog.feedback.text()
    assert tab.editor.metadata().values["/Title"] == "Edited document title"
    assert tab.editor.dirty and window.a_undo.isEnabled()
    assert tab.page == 1
    output = tmp_path / "with_metadata.pdf"
    tab.editor.save_as(output)
    assert PdfReader(output).metadata.title == "Edited document title"
    assert not tab.editor.dirty
    dialog.close()
    dialog.deleteLater()
    window.undo()
    assert tab.editor.metadata().values["/Title"] == ""
    window.redo()
    assert tab.editor.metadata().values["/Title"] == "Edited document title"


def test_properties_action_and_banner_open_same_inspector(window):
    seen = []
    def close_inspector():
        dialogs = [w for w in QApplication.topLevelWidgets() if isinstance(w, PropertiesDialog) and w.isVisible()]
        assert len(dialogs) == 1
        seen.append(dialogs[0].tabs.count())
        dialogs[0].reject()
    QTimer.singleShot(0, close_inspector)
    window.a_info.trigger()
    QTimer.singleShot(0, close_inspector)
    window.current.notice.show_details()
    assert seen == [5, 5]
    assert not window.current.editor.dirty


def test_metadata_dialog_draft_close_leaves_pdf_unchanged(window):
    tab = window.current
    revision = tab.editor.path
    dialog = PropertiesDialog(tab.editor, lambda values: tab.apply_edit("update_metadata", values), window)
    dialog.fields["/Title"].setText("Unapplied draft")
    assert dialog.apply_button.isEnabled()
    dialog.reject()
    dialog.deleteLater()
    assert tab.editor.path == revision and not tab.editor.dirty


def test_encrypted_metadata_inspector_is_readonly(window, tmp_path, monkeypatch):
    writer = PdfWriter()
    writer.add_blank_page(400, 600)
    writer.add_metadata({"/Title": "Protected document"})
    writer.encrypt("secret")
    source = tmp_path / "protected.pdf"
    writer.write(source)
    monkeypatch.setattr(QInputDialog, "getText", lambda *_args, **_kwargs: ("secret", True))
    window.open_pdf(source)
    tab = window.current
    assert tab.editor.encrypted and window.a_info.isEnabled()
    assert tab.notice.isVisible() and "sola lettura" in tab.notice.text.text()
    wait_for(lambda: 0 in tab.viewer.vectors.cache, timeout=10000)
    assert tab.viewer.vectors.password == "secret"
    dialog = PropertiesDialog(tab.editor, lambda values: tab.apply_edit("update_metadata", values), window)
    assert dialog.fields["/Title"].text() == "Protected document"
    assert all(field.isReadOnly() for field in dialog.fields.values())
    assert not dialog.apply_button.isEnabled()
    assert "sola lettura" in dialog.warning.text()
    dialog.reject()
    dialog.deleteLater()


@pytest.mark.parametrize("algorithm", ["RC4-128", "AES-128", "AES-256"])
def test_modify_restricted_pdf_opens_without_password_prompt(window, text_pdf, tmp_path, monkeypatch, algorithm):
    from pypdf.constants import UserAccessPermissions
    source = tmp_path / "modify-restricted.pdf"
    writer = PdfWriter(clone_from=text_pdf)
    writer.add_metadata({"/Title": "Restricted document"})
    writer.encrypt("", owner_password="owner-secret", algorithm=algorithm,
                   permissions_flag=UserAccessPermissions.PRINT | UserAccessPermissions.EXTRACT)
    writer.write(source)
    original = source.read_bytes()
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: pytest.fail("A reading password was requested"))

    window.open_pdf(source)
    tab = window.current
    assert tab.editor.source == source.resolve()
    assert tab.editor.encrypted and not tab.viewer.editable
    assert tab.pdf.status() == QPdfDocument.Status.Ready and tab.count == 3
    assert "Alpha beta PDF page 1" in tab.pdf.getAllText(0).text()
    wait_for(lambda: 0 in tab.viewer.vectors.cache, timeout=10000)
    assert not tab.viewer.vectors.password
    assert window.document_state.text() == "Sola lettura"
    assert tab.notice.isVisible()
    assert "Modifiche disabilitate" in tab.notice.text.text()
    assert "sola lettura" in tab.notice.text.text()
    assert "salvataggio sono disabilitati" in tab.notice.toolTip()
    assert tab.notice.height() <= 64
    for action in (window.a_save_current, window.a_save, window.a_rotate, window.a_note,
                   window.a_export_docx, window.a_ocr):
        assert not action.isEnabled()
    assert window.a_info.isEnabled()
    assert tab.editor.metadata().values["/Title"] == "Restricted document"
    assert not tab.editor.metadata().editable
    assert source.read_bytes() == original


def test_page_number_remains_visible_and_accepts_keyboard_navigation(window):
    assert window.page_spin.width() >= 32
    window.page_spin.setFocus()
    window.page_spin.selectAll()
    QTest.keyClicks(window.page_spin, "3")
    QTest.keyClick(window.page_spin, Qt.Key.Key_Return)
    assert window.current.page == 2 and window.page_spin.value() == 3


def test_page_number_keeps_multi_digit_draft_during_document_updates(window, tmp_path):
    writer = PdfWriter()
    for _ in range(95):
        writer.add_blank_page(width=400, height=600)
    source = tmp_path / "many-pages.pdf"
    writer.write(source)
    window.open_pdf(source)
    QTest.qWait(150)
    tab = window.current
    window.page_spin.setFocus()
    window.page_spin.selectAll()
    QTest.keyClicks(window.page_spin, "8")
    window._update_state()
    QTest.qWait(50)
    assert window.page_spin.text() == "8"
    QTest.keyClicks(window.page_spin, "0")
    assert tab.page == 0
    QTest.keyClick(window.page_spin, Qt.Key.Key_Return)
    wait_for(lambda: tab.page == 79)
    QTest.qWait(150)
    rect = tab.viewer.page_rects()[79].translated(
        -tab.viewer.horizontalScrollBar().value(), -tab.viewer.verticalScrollBar().value())
    assert rect.intersects(tab.viewer.viewport().rect())
    assert tab.thumbs.currentRow() == 79
    assert window.page_spin.text() == "80"


@pytest.mark.parametrize("number", ["0", "4", "", "abc", "-1"])
def test_page_number_invalid_input_does_not_navigate(window, number):
    tab = window.current
    tab.jump_to_page(1)
    QTest.qWait(50)
    window.page_spin.setFocus()
    window.page_spin.selectAll()
    QTest.keyClick(window.page_spin, Qt.Key.Key_Backspace)
    QTest.keyClicks(window.page_spin, number)
    QTest.keyClick(window.page_spin, Qt.Key.Key_Return)
    assert tab.page == 1
    assert window.page_spin.text() == "2"
    assert "tra 1 e 3" in window.statusBar().currentMessage()


def test_page_number_focus_confirmation_escape_and_same_page(window):
    requested = []
    window.page_spin.page_requested.connect(requested.append)
    window.page_spin.setFocus()
    window.page_spin.selectAll()
    QTest.keyClicks(window.page_spin, "3")
    window.zoom_combo.setFocus()
    assert window.current.page == 2
    assert requested == [3]
    window.page_spin.setFocus()
    window.page_spin.selectAll()
    QTest.keyClicks(window.page_spin, "1")
    QTest.keyClick(window.page_spin, Qt.Key.Key_Escape)
    window.zoom_combo.setFocus()
    assert requested == [3] and window.page_spin.text() == "3"
    window.page_spin.setFocus()
    QTest.keyClick(window.page_spin, Qt.Key.Key_Return)
    assert requested == [3, 3]


def test_page_number_draft_is_reset_on_tab_change(window, text_pdf, tmp_path):
    first = window.current
    another = tmp_path / "another.pdf"
    another.write_bytes(text_pdf.read_bytes())
    window.open_pdf(another)
    second = window.current
    window.page_spin.setFocus()
    window.page_spin.selectAll()
    QTest.keyClicks(window.page_spin, "3")
    window.tabs.setCurrentWidget(first)
    assert window.page_spin.text() == "1"
    QTest.keyClick(window.page_spin, Qt.Key.Key_Return)
    assert first.page == 0 and second.page == 0


@pytest.fixture
def isolated_theme(app):
    settings = QSettings(QSettings.defaultFormat(), QSettings.Scope.UserScope, "APEX", "PDF Studio")
    previous = settings.value("ui/theme")
    settings.remove("ui/theme")
    apply_theme(app)
    yield settings
    if previous is None:
        settings.remove("ui/theme")
    else:
        settings.setValue("ui/theme", previous)
    settings.sync()
    apply_theme(app)


def test_theme_defaults_to_blue_with_invalid_setting(app, isolated_theme):
    for value in (None, "missing-theme"):
        if value:
            isolated_theme.setValue("ui/theme", value)
        widget = MainWindow()
        assert widget.theme_name == DEFAULT_THEME == "blue"
        assert theme_color("accent") == THEMES["blue"].accent
        assert widget.a_settings.isEnabled() and widget.current is None
        widget.close()
        widget.deleteLater()


@pytest.mark.parametrize("name", list(THEMES))
def test_theme_settings_apply_and_persist_without_restart(window, app, isolated_theme, name):
    tab = window.current
    tab.set_zoom(125)
    tab.jump_to_page(1)
    window.focus_search()
    window.search.setText("PDF")
    wait_for(lambda: tab.search_model.count() == 3)
    tab.jump_to_page(1)
    QTest.qWait(100)
    state = (tab.editor.path, tab.editor.dirty, tab.page, tab.viewer.zoomFactor(), tab._query)
    old_icon = window.a_open.icon().cacheKey()
    dialog = SettingsDialog(window.theme_name, window)
    dialog.theme_applied.connect(window._apply_color_theme)
    dialog.show()
    dialog.theme_combo.setCurrentIndex(dialog.theme_combo.findData(name))
    QTest.mouseClick(dialog.apply_button, Qt.MouseButton.LeftButton)
    QTest.qWait(100)
    assert window.theme_name == name
    assert app.palette().color(QPalette.ColorRole.Highlight).name() == THEMES[name].accent
    assert theme_colors(name)["notice"] in app.styleSheet()
    assert tab.viewer.viewport().grab().toImage().pixelColor(0, 0).name() == theme_colors(name)["canvas"]
    assert window.a_open.icon().cacheKey() != old_icon
    assert isolated_theme.value("ui/theme") == name
    assert (tab.editor.path, tab.editor.dirty, tab.page, tab.viewer.zoomFactor(), tab._query) == state
    assert tab.thumbs.currentRow() == tab.page
    dialog.reject()
    dialog.deleteLater()
    restarted = MainWindow()
    assert restarted.theme_name == name
    restarted.close()
    restarted.deleteLater()


def test_theme_settings_cancel_preview_and_restore_blue(window, app, isolated_theme):
    dialog = SettingsDialog(window.theme_name, window)
    dialog.theme_applied.connect(window._apply_color_theme)
    dialog.show()
    original_style = app.styleSheet()
    dialog.theme_combo.setCurrentIndex(dialog.theme_combo.findData("green"))
    assert THEMES["green"].accent in dialog.preview.styleSheet()
    dialog.reject()
    assert app.styleSheet() == original_style
    assert isolated_theme.value("ui/theme") is None
    dialog.show()
    dialog.apply_selection()
    assert window.theme_name == "green"
    QTest.mouseClick(dialog.reset_button, Qt.MouseButton.LeftButton)
    assert dialog.theme_combo.currentData() == "blue"
    assert window.theme_name == "green"
    dialog.accept()
    assert window.theme_name == "blue"
    assert isolated_theme.value("ui/theme") == "blue"
    dialog.deleteLater()


def test_security_rejects_active_pdf_before_qt_load(window, tmp_path, monkeypatch):
    from pypdf.generic import TextStringObject
    writer = PdfWriter()
    writer.add_blank_page(400, 600)
    writer.root_object[NameObject("/OpenAction")] = writer._add_object(DictionaryObject({
        NameObject("/S"): NameObject("/JavaScript"), NameObject("/JS"): TextStringObject("synthetic test")}))
    source = tmp_path / "active.pdf"
    writer.write(source)
    errors = []
    monkeypatch.setattr(QMessageBox, "critical", lambda *_args: errors.append(_args[2]))
    monkeypatch.setattr("ape_pdf_studio.ui.document_tab.QPdfDocument", lambda *_a: pytest.fail("active PDF reached Qt"))
    previous = window.current
    window.open_pdf(source)
    assert window.tabs.count() == 1 and window.current is previous
    assert len(errors) == 1 and "JavaScript" in errors[0]


def test_security_details_and_attachment_notice_are_visible(window, tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(400, 600)
    writer.add_attachment("sample.txt", b"synthetic attachment")
    path = tmp_path / "attached.pdf"
    writer.write(path)
    window.open_pdf(path)
    QTest.qWait(100)
    tab = window.current
    assert tab.notice.isVisible()
    assert tab.notice.property("securityWarning")
    assert "Allegati" in tab.notice.text.text()
    dialog = PropertiesDialog(tab.editor, lambda values: tab.apply_edit("update_metadata", values), window)
    text = dialog.security_details.toPlainText()
    assert "Controllo preventivo attivo" in text
    assert "sandbox" in text and "antivirus" in text
    assert "Allegati" in text
    assert not tab.editor.dirty
    dialog.deleteLater()


def test_large_page_zoom_has_bounded_raster_requests(window, tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(12000, 12000)
    path = tmp_path / "large-sheet.pdf"
    writer.write(path)
    window.open_pdf(path)
    QTest.qWait(100)
    viewer = window.current.viewer
    viewer.set_vector_enabled(False)
    window.current.set_zoom(500)
    viewer.viewport().grab()
    keys = list(viewer._images) + list(viewer._requests.values())
    assert keys
    assert all(tile.width * tile.height <= 1024 ** 2 for tile in keys)
    assert max(tile.full_width for tile in keys) >= 80_000
    assert len(viewer._requests) <= 2


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_tiles_match_full_render_on_rotated_cropped_and_annotated_page(window, text_pdf, tmp_path, rotation):
    import numpy as np
    from pypdf.generic import RectangleObject, ArrayObject, FloatObject
    from pypdf.annotations import Highlight
    from PySide6.QtCore import QRect, QSize
    writer = PdfWriter()
    page = writer.add_page(PdfReader(text_pdf).pages[0])
    page.cropbox = RectangleObject([20, 30, 380, 580])
    page.rotate(rotation)
    writer.add_annotation(0, Highlight(rect=(48, 489, 125, 515),
                                     quad_points=ArrayObject([FloatObject(x) for x in (48, 515, 125, 515, 48, 489, 125, 489)])))
    path = tmp_path / "tile-test.pdf"
    writer.write(path)
    window.open_pdf(path)
    tab = window.current
    tab.set_zoom(250)
    viewer = tab.viewer
    QTest.qWait(50)
    viewer.set_vector_enabled(False)
    viewer.horizontalScrollBar().setValue(viewer.horizontalScrollBar().maximum() // 2)
    rect = viewer.page_rects()[0]
    viewer.verticalScrollBar().setValue(int(rect.top() + 1024 / viewer.devicePixelRatioF() - viewer.viewport().height() / 2))
    viewer.viewport().grab()
    wait_for(lambda: viewer._wanted and all(tile in viewer._images for tile in viewer._wanted))
    options = QPdfDocumentRenderOptions()
    options.setRenderFlags(QPdfDocumentRenderOptions.RenderFlag.Annotations)
    first = next(iter(viewer._wanted))
    reference = tab.pdf.render(0, QSize(first.full_width, first.full_height), options)
    assert not reference.isNull()
    for tile in viewer._wanted:
        expected = reference.copy(QRect(tile.left, tile.top, tile.width, tile.height))
        actual = viewer._images[tile]
        a = np.frombuffer(actual.constBits(), dtype=np.uint8).astype(np.int16)
        b = np.frombuffer(expected.constBits(), dtype=np.uint8).astype(np.int16)
        assert actual.size() == expected.size()
        assert np.mean(np.abs(a - b)) < .15
    assert len(viewer._wanted) > 1


def test_fast_zoom_keeps_requests_bounded_and_renders_the_final_view(window):
    from ape_pdf_studio.ui.page_tiles import CACHE_BYTES
    viewer = window.current.viewer
    viewer.set_vector_enabled(False)
    QTest.qWait(50)
    for zoom in (300, 500, 100, 250, 400):
        window.current.set_zoom(zoom)
        viewer.viewport().grab()
        assert len(viewer._requests) <= 2
    wait_for(lambda: viewer._wanted and all(tile in viewer._images for tile in viewer._wanted))
    rect = viewer.page_rects()[0]
    expected_width = round(rect.width() * viewer.devicePixelRatioF())
    assert all(tile.full_width == expected_width for tile in viewer._wanted)
    assert viewer._cache_bytes <= CACHE_BYTES


def test_revision_change_clears_tiles_and_ignores_old_renderer_results(window):
    viewer = window.current.viewer
    viewer.set_vector_enabled(False)
    wait_for(lambda: bool(viewer._images))
    previous_renderer = viewer._renderer
    epoch = viewer._render_epoch
    window._edit("rotate", 0, 90)
    assert viewer._render_epoch > epoch
    assert viewer._renderer is not previous_renderer
    before = dict(viewer._images)
    # A retired renderer may still have a completion signal queued.
    previous_renderer.pageRendered.emit(0, QImage(4, 4, QImage.Format.Format_ARGB32).size(),
                                        QImage(4, 4, QImage.Format.Format_ARGB32), QPdfDocumentRenderOptions(), 1)
    assert dict(viewer._images) == before
    wait_for(lambda: viewer._wanted and all(tile in viewer._images for tile in viewer._wanted))


def test_vector_zoom_reuses_glyph_paths_and_keeps_selection_and_search(window):
    tab = window.current
    viewer = tab.viewer
    wait_for(lambda: 0 in viewer.vectors.cache, timeout=10000)
    page = viewer.vectors.cache[0]
    assert window.render_state.text() == "Vettoriale"
    for zoom in (100, 200, 500, 125):
        tab.set_zoom(zoom)
        viewer.viewport().grab()
        assert viewer.vectors.cache[0] is page
        assert not any(tile.page == 0 for tile in viewer._wanted)
    viewer.select_page_text(0)
    viewer.copy_selection()
    assert "Alpha beta PDF" in QApplication.clipboard().text()
    window.search.setText("Alpha")
    wait_for(lambda: tab.search_model.count() == 3)
    assert viewer.currentSearchResultIndex() == 0


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_vector_clips_match_pdf_renderer_after_crop_and_rotation(window, tmp_path, rotation):
    import numpy as np
    from pypdf.generic import RectangleObject
    from PySide6.QtCore import QRectF, QSize
    from PySide6.QtGui import QPainter, QTransform
    writer = PdfWriter()
    page = writer.add_blank_page(100, 200)
    page.cropbox = RectangleObject([10, 20, 90, 180])
    page.rotate(rotation)
    stream = DecodedStreamObject()
    stream.set_data(b"0 0 1 rg 0 0 100 200 re f q 20 30 40 70 re W n 1 0 0 rg 0 0 100 200 re f Q")
    page[NameObject("/Contents")] = writer._add_object(stream)
    source = tmp_path / "clipped.pdf"
    writer.write(source)
    window.open_pdf(source)
    tab = window.current
    wait_for(lambda: 0 in tab.viewer.vectors.cache, timeout=10000)
    size = (tab.pdf.pagePointSize(0) * 4).toSize()
    image = QImage(size, QImage.Format.Format_ARGB32_Premultiplied)
    image.fill(Qt.GlobalColor.white)
    painter = QPainter(image)
    painter.setRenderHint(QPainter.RenderHint.Antialiasing)
    transform = QTransform.fromScale(4, 4)
    tab.viewer.vectors.cache[0].paint(painter, transform, QRectF(image.rect()))
    painter.end()
    reference = tab.pdf.render(0, QSize(size))
    a = np.frombuffer(image.constBits(), dtype=np.uint8).reshape(size.height(), size.width(), 4)
    b = np.frombuffer(reference.constBits(), dtype=np.uint8).reshape(size.height(), size.width(), 4)
    assert np.mean(np.abs(a.astype(np.int16) - b.astype(np.int16))) < 1
    assert any(layer[2] for layer in tab.viewer.vectors.cache[0].layers)


def test_vector_revision_reload_undo_and_manual_compatibility(window):
    tab = window.current
    vectors = tab.viewer.vectors
    wait_for(lambda: 0 in vectors.cache, timeout=10000)
    original = vectors.cache[0]
    window._edit("rotate", 0, 90)
    assert original not in vectors.cache.values()
    wait_for(lambda: 0 in vectors.cache, timeout=10000)
    assert vectors.cache[0].renderer.viewBoxF().width() == 600
    window.undo()
    wait_for(lambda: 0 in vectors.cache, timeout=10000)
    assert vectors.cache[0].renderer.viewBoxF().width() == 400
    window.a_vector.trigger()
    assert not vectors.enabled and not vectors.cache and vectors.process is None
    assert window.render_state.text() == "Compatibilità"
    window.a_vector.trigger()
    wait_for(lambda: 0 in vectors.cache, timeout=10000)
    assert window.render_state.text() == "Vettoriale"


def test_embedded_image_uses_complete_compatibility_page(window, tmp_path):
    import pymupdf
    image = QImage(20, 20, QImage.Format.Format_RGB32)
    image.fill(Qt.GlobalColor.red)
    png = tmp_path / "source.png"
    image.save(str(png))
    pdf = pymupdf.open()
    page = pdf.new_page(width=400, height=600)
    page.insert_image(pymupdf.Rect(40, 50, 200, 210), filename=str(png))
    page.insert_text((40, 240), "Image and text")
    source = tmp_path / "mixed.pdf"
    pdf.save(source)
    pdf.close()
    window.open_pdf(source)
    viewer = window.current.viewer
    wait_for(lambda: 0 in viewer.vectors.failed, timeout=10000)
    assert 0 not in viewer.vectors.cache
    assert window.render_state.text() == "Compatibilità"
    viewer.viewport().grab()
    wait_for(lambda: viewer._wanted and all(tile in viewer._images for tile in viewer._wanted))


def test_highlight_switches_to_compatibility_and_retains_annotation(window):
    tab = window.current
    viewer = tab.viewer
    wait_for(lambda: 0 in viewer.vectors.cache, timeout=10000)
    viewer.select_page_text(0)
    window.markup("highlight")
    wait_for(lambda: 0 in viewer.vectors.failed, timeout=10000)
    assert window.render_state.text() == "Compatibilità"
    viewer.viewport().grab()
    wait_for(lambda: viewer._wanted and all(tile in viewer._images for tile in viewer._wanted))
    assert PdfReader(tab.editor.path).pages[0]["/Annots"][0].get_object()["/Subtype"] == "/Highlight"
    yellow = 0
    for tile, image in viewer._images.items():
        if tile.page != 0:
            continue
        for y in range(0, image.height(), 3):
            for x in range(0, image.width(), 3):
                color = image.pixelColor(x, y)
                yellow += color.red() > 200 and color.green() > 200 and color.blue() < 100
    assert yellow > 10


def test_close_stops_vector_worker_before_revision_cleanup(window):
    tab = window.current
    vectors = tab.viewer.vectors
    # Force a fresh job, then close while it is queued/running.
    vectors.reset(tab.editor.path)
    vectors.request([0])
    process = vectors.process
    assert process is not None
    revision = tab.editor.path
    tab.close_document()
    assert vectors.process is None and not vectors.cache and vectors.source is None
    assert not revision.exists()
    from PySide6.QtCore import QProcess
    assert process.state() == QProcess.ProcessState.NotRunning


def test_completion_toast_keeps_focus_repositions_and_dismisses(window, tmp_path):
    window.search.setFocus()
    focus = QApplication.focusWidget()
    path = tmp_path / ("<documento>" + "lungo" * 40 + ".pdf")
    window.toast.show_completion("PDF salvato", path)
    assert QApplication.focusWidget() is focus
    assert window.toast.isVisible()
    assert window.toast.filename.textFormat() == Qt.TextFormat.PlainText
    assert window.toast.filename.toolTip() == str(path)
    assert "…" in window.toast.filename.text()
    window.resize(950, 650)
    QTest.qWait(30)
    assert window.toast.geometry().right() < window.width()
    assert window.toast.geometry().bottom() < window.height() - window.statusBar().height()
    QTest.mouseClick(window.toast.close_button, Qt.MouseButton.LeftButton)
    assert not window.toast.isVisible() and not window.toast.timer.isActive()
    window.toast.show_completion("Immagine esportata", tmp_path / "page.png")
    window.toast.timer.start(40)
    wait_for(lambda: not window.toast.isVisible())


def test_cancelled_or_failed_save_does_not_show_success_toast(window, tmp_path, monkeypatch):
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: ("", ""))
    window.save_as()
    assert not window.toast.isVisible()
    path = tmp_path / "failure.pdf"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(path), ""))
    errors = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *a: errors.append(a[2]))
    def fail_save(destination):
        raise PDFOperationError("Salvataggio non riuscito")
    monkeypatch.setattr(window.current.editor, "save_as", fail_save)
    window.save_as()
    assert errors and not path.exists() and not window.toast.isVisible()


def test_extracted_pages_show_completion_toast_after_file_is_written(window, tmp_path, monkeypatch):
    output = tmp_path / "pages.pdf"
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: ("2", True))
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(output), ""))
    window.extract_pages()
    assert len(PdfReader(output).pages) == 1
    assert window.toast.isVisible() and window.toast.title.text() == "Pagine PDF esportate"
    assert window.toast.filename.text() == output.name


def test_context_note_uses_clicked_position_and_opens_edits_deletes_and_undoes(window, monkeypatch):
    tab = window.current
    point = QPointF(150, 220)
    def create(dialog):
        dialog.text.setPlainText("Nota precisa")
        dialog._save()
        return dialog.result()
    monkeypatch.setattr(NoteDialog, "exec", create)
    menu = tab.viewer.make_context_menu(0, point)
    next(a for a in menu.actions() if a.text() == "Aggiungi nota qui…").trigger()
    entry = tab.viewer.notes[0][0]
    assert entry["rect"] == (150, 220, 174, 244) and entry["text"] == "Nota precisa"
    def edit(dialog):
        assert dialog.text.toPlainText() == "Nota precisa"
        dialog.text.setPlainText("Nota corretta")
        dialog._save()
        return dialog.result()
    monkeypatch.setattr(NoteDialog, "exec", edit)
    position = tab.viewer.page_transform(0).map(QPointF(162, 232)).toPoint()
    QTest.mouseClick(tab.viewer.viewport(), Qt.MouseButton.LeftButton, pos=position)
    assert tab.viewer.notes[0][0]["text"] == "Nota corretta"
    def delete(dialog):
        dialog.delete_button.click()
        return dialog.result()
    monkeypatch.setattr(NoteDialog, "exec", delete)
    position = tab.viewer.page_transform(0).map(QPointF(162, 232)).toPoint()
    QTest.mouseClick(tab.viewer.viewport(), Qt.MouseButton.LeftButton, pos=position)
    assert not tab.viewer.notes
    window.undo()
    assert tab.viewer.notes[0][0]["text"] == "Nota corretta"
    window.undo()
    assert tab.viewer.notes[0][0]["text"] == "Nota precisa"
    menu.deleteLater()


def test_toolbar_note_placement_is_cancelled_by_escape_and_creates_at_click(window, monkeypatch):
    view = window.current.viewer
    window.a_note.trigger()
    assert view._placement == "note"
    QTest.keyClick(view, Qt.Key.Key_Escape)
    assert view._placement is None and not window.current.editor.dirty
    def save(dialog):
        dialog.text.setPlainText("Nota dal pulsante")
        dialog._save()
        return dialog.result()
    monkeypatch.setattr(NoteDialog, "exec", save)
    window.a_note.trigger()
    position = view.page_transform(0).map(QPointF(120, 200)).toPoint()
    QTest.mouseClick(view.viewport(), Qt.MouseButton.LeftButton, pos=position)
    note = view.notes[0][0]
    assert note["rect"][:2] == pytest.approx((120, 200), abs=1)
    assert view._placement is None


def test_direct_text_context_edit_updates_search_save_and_undo(window, tmp_path, monkeypatch):
    tab = window.current
    bounds = tab.pdf.getAllText(0).bounds()[0].boundingRect()
    point = QPointF(bounds.left() + 3, bounds.center().y())
    def edit(dialog):
        assert "Alpha beta PDF" in dialog.text.text()
        dialog.text.setText("Testo corretto")
        dialog._save()
        assert not dialog.feedback.isVisible()
        return dialog.result()
    monkeypatch.setattr(TextEditDialog, "exec", edit)
    menu = tab.viewer.make_context_menu(0, point)
    action = next(a for a in menu.actions() if a.text() == "Modifica testo…")
    assert action.isEnabled()
    action.trigger()
    assert "Testo corretto" in tab.pdf.getAllText(0).text()
    assert "Alpha beta PDF" not in tab.pdf.getAllText(0).text()
    assert "Second line" in tab.pdf.getAllText(0).text()
    assert tab.viewer.edit_target is None
    window.search.setText("Testo corretto")
    wait_for(lambda: tab.search_model.count() == 1)
    output = tmp_path / "new-text.pdf"
    tab.editor.save_as(output)
    assert "Testo corretto" in PdfReader(output).pages[0].extract_text()
    window.undo()
    assert "Alpha beta PDF" in tab.pdf.getAllText(0).text()
    window.redo()
    assert "Testo corretto" in tab.pdf.getAllText(0).text()
    menu.deleteLater()


def test_text_dialog_keeps_invalid_edit_open_and_encrypted_commands_disabled(window, tmp_path, monkeypatch):
    tab = window.current
    target = tab.editor.text_at(0, (60, 95))
    original = tab.editor.path
    dialog = TextEditDialog(tab.editor.path, target, lambda text, size: tab.apply_edit("replace_text", target, text, size), window)
    dialog.show()
    dialog.text.setText("Testo troppo lungo " * 50)
    dialog.save_button.click()
    assert dialog.isVisible() and dialog.feedback.isVisible() and tab.editor.path == original
    dialog.reject()
    dialog.deleteLater()
    writer = PdfWriter()
    writer.append(PdfReader(tab.editor.path))
    writer.encrypt("secret")
    source = tmp_path / "readonly.pdf"
    writer.write(source)
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: ("secret", True))
    window.open_pdf(source)
    assert not window.a_edit_text.isEnabled() and not window.a_note.isEnabled()
    menu = window.current.viewer.make_context_menu(0, QPointF(60, 95))
    for title in ("Aggiungi nota qui…", "Modifica testo…"):
        assert not next(a for a in menu.actions() if a.text() == title).isEnabled()
    menu.deleteLater()
