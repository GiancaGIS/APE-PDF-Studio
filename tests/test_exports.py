"""Conversion keeps every page, editable text and scan images, atomically."""
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
import pymupdf

from ape_pdf_studio.core.docx_export import atomic_copy, convert_pdf_to_docx
from ape_pdf_studio.core.editor import PDFOperationError


@pytest.fixture
def mixed_pdf(tmp_path):
    pdf = pymupdf.open()
    first = pdf.new_page(width=400, height=600)
    first.insert_text((40, 80), "Editable first page", fontsize=16)
    first.insert_text((40, 110), "Text survives conversion.")
    scan = pdf.new_page(width=400, height=600)
    pixels = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 300), False)
    pixels.clear_with(170)
    scan.insert_image(scan.rect, pixmap=pixels)
    pdf.new_page(width=400, height=600)  # Blank page between scan and last page.
    last = pdf.new_page(width=600, height=400)
    last.insert_text((40, 80), "Editable last page", fontsize=16)
    path = tmp_path / "mixed.pdf"
    pdf.save(path)
    pdf.close()
    return path


def test_word_keeps_text_scans_blank_pages_dimensions_and_order(mixed_pdf, tmp_path):
    original = mixed_pdf.read_bytes()
    events = []
    output = tmp_path / "document.docx"
    result = convert_pdf_to_docx(mixed_pdf, output, lambda *args: events.append(args))
    assert result == {"pages": 4, "image_pages": [2, 3]}
    document = Document(output)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Editable first page" in text
    assert "Editable last page" in text
    assert text.index("first page") < text.index("last page")
    assert len(document.sections) == 4
    assert len(document.inline_shapes) == 2
    assert [(round(s.page_width.pt), round(s.page_height.pt)) for s in document.sections] == [(400, 600)] * 3 + [(600, 400)]
    assert events[0][:2] == (0, 4) and events[-1][:2] == (4, 4)
    assert mixed_pdf.read_bytes() == original
    with ZipFile(output) as package:
        assert package.testzip() is None
        assert len([n for n in package.namelist() if n.startswith("word/media/")]) == 2


def test_all_blank_pages_are_retained(tmp_path):
    source = tmp_path / "blank.pdf"
    with pymupdf.open() as pdf:
        for _ in range(3):
            pdf.new_page()
        pdf.save(source)
    output = tmp_path / "blank.docx"
    assert convert_pdf_to_docx(source, output)["image_pages"] == [1, 2, 3]
    assert len(Document(output).sections) == 3


def test_conversion_failure_preserves_existing_output(mixed_pdf, tmp_path, monkeypatch):
    from pdf2docx.page.Page import Page

    def fail(*_args, **_kwargs):
        raise ValueError("layout failure")

    output = tmp_path / "existing.docx"
    output.write_bytes(b"original output")
    monkeypatch.setattr(Page, "parse", fail)
    with pytest.raises(PDFOperationError, match="layout failure"):
        convert_pdf_to_docx(mixed_pdf, output)
    assert output.read_bytes() == b"original output"
    assert not list(tmp_path.glob(".apex_export_*"))


def test_word_export_rejects_pdf_destination_and_encrypted_input(mixed_pdf, tmp_path):
    original = mixed_pdf.read_bytes()
    with pytest.raises(PDFOperationError, match=".docx"):
        convert_pdf_to_docx(mixed_pdf, mixed_pdf)
    assert mixed_pdf.read_bytes() == original
    encrypted = tmp_path / "encrypted.pdf"
    with pymupdf.open(mixed_pdf) as pdf:
        pdf.save(encrypted, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="secret", user_pw="secret")
    with pytest.raises(PDFOperationError, match="protetti"):
        convert_pdf_to_docx(encrypted, tmp_path / "encrypted.docx")


def test_atomic_export_cleans_temporary_on_failed_commit(tmp_path, monkeypatch):
    import ape_pdf_studio.core.docx_export as module

    source = tmp_path / "completed.docx"
    source.write_bytes(b"new")
    output = tmp_path / "destination.docx"
    output.write_bytes(b"old")
    def fail(*_args):
        raise PermissionError("locked")
    monkeypatch.setattr(module.os, "replace", fail)
    with pytest.raises(PermissionError):
        atomic_copy(source, output)
    assert output.read_bytes() == b"old"
    assert not list(tmp_path.glob(".apex_export_*"))
